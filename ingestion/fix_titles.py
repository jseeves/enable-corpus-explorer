"""
fix_titles.py — extract real publication titles from PDFs and patch umap.json

For each document:
  1. Try the PDF's embedded metadata title (fast, clean when present)
  2. Otherwise extract the first page text and ask Claude to identify the title
  3. Write the result back to app/public/umap.json

Usage:
    cd ingestion
    python fix_titles.py [--dry-run]

Requires:
    pip install pypdf anthropic openpyxl python-docx
"""

import argparse
import json
import os
import sys
import time

import openpyxl
import pypdf

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL_PATH = os.path.join(REPO_ROOT, "data", "metadata", "Enable_Stocktake_v6.xlsx")
PDF_DIR    = os.path.join(REPO_ROOT, "data", "raw_pdfs")
UMAP_PATH  = os.path.join(REPO_ROOT, "app", "public", "umap.json")


# ── helpers ──────────────────────────────────────────────────────────────────

def load_excel_map():
    """Return {resource_id: file_name} from the Corpus_Classification sheet."""
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    ws = wb["Corpus_Classification"]
    rows = ws.iter_rows(values_only=True)
    headers = next(rows)
    rid_col  = headers.index("resource_id")
    file_col = headers.index("file_name")
    result = {}
    for row in rows:
        rid  = row[rid_col]
        fname = row[file_col]
        if rid and fname:
            result[str(rid).strip()] = str(fname).strip()
    return result


def find_pdf(file_name: str):
    """Return the full path to the file if it exists in PDF_DIR."""
    candidate = os.path.join(PDF_DIR, file_name)
    if os.path.exists(candidate):
        return candidate
    # case-insensitive fallback
    lower = file_name.lower()
    for f in os.listdir(PDF_DIR):
        if f.lower() == lower:
            return os.path.join(PDF_DIR, f)
    return None


def extract_pdf_meta_title(path: str) :
    """Return the PDF's embedded Title metadata, or None."""
    try:
        with open(path, "rb") as f:
            reader = pypdf.PdfReader(f)
            meta = reader.metadata
            if meta and meta.title and len(meta.title.strip()) > 4:
                # Strip common suffixes like " | World Resources Institute"
                title = meta.title.strip()
                # Strip site-name suffixes appended by web-to-PDF tools
                for suffix in [
                    " | World Resources Institute",
                    " - World Resources Institute",
                    " _ World Resources Institute",
                    " | WRI",
                ]:
                    if title.endswith(suffix):
                        title = title[: -len(suffix)].strip()
                # Strip any remaining "| <site name>" tail
                if " | " in title:
                    title = title[: title.rfind(" | ")].strip()
                return title
    except Exception:
        pass
    return None


def extract_first_page_text(path: str) -> str:
    """Return the first ~600 chars of text from the first page of a PDF."""
    try:
        if path.lower().endswith(".docx"):
            import docx
            doc = docx.Document(path)
            text = "\n".join(p.text for p in doc.paragraphs[:20])
            return text[:600]
        with open(path, "rb") as f:
            reader = pypdf.PdfReader(f)
            if not reader.pages:
                return ""
            return reader.pages[0].extract_text()[:600]
    except Exception as e:
        print(f"    [warn] could not extract text: {e}")
        return ""


def claude_extract_title(resource_id: str, file_name: str, first_page: str) -> str:
    """Ask Claude to extract the actual publication title from the first-page text."""
    import anthropic
    client = anthropic.Anthropic(api_key=os.environ["ANTHROPIC_API_KEY"])

    prompt = f"""I have a WRI document with filename: {file_name}

Here is the text from its first page:
---
{first_page}
---

What is the actual publication title of this document? Return ONLY the title — no explanation, no quotes, no punctuation beyond what appears in the title itself. If the title is in ALL CAPS in the text, convert it to Title Case."""

    msg = client.messages.create(
        model=os.environ.get("CLAUDE_MODEL", "claude-sonnet-4-5"),
        max_tokens=100,
        messages=[{"role": "user", "content": prompt}],
    )
    return msg.content[0].text.strip().strip('"').strip("'")


# ── main ─────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true", help="Print titles without writing")
    args = parser.parse_args()

    # Load current umap.json
    with open(UMAP_PATH) as f:
        umap = json.load(f)
    docs = umap.get("docs", umap) if isinstance(umap, dict) else umap

    # Load Excel resource_id → file_name mapping
    rid_to_file = load_excel_map()
    print(f"Loaded {len(rid_to_file)} entries from Excel")

    results: dict[str, str] = {}  # resource_id → new title
    skipped = 0

    for doc in docs:
        rid   = doc["resource_id"]
        old_title = doc.get("title", "")
        fname = rid_to_file.get(rid)

        if not fname:
            print(f"  [{rid}] NOT IN EXCEL — keeping: {old_title!r}")
            skipped += 1
            continue

        path = find_pdf(fname)
        if not path:
            print(f"  [{rid}] FILE NOT FOUND: {fname!r} — keeping: {old_title!r}")
            skipped += 1
            continue

        # Step 1: PDF metadata title
        title = extract_pdf_meta_title(path)
        method = "metadata"

        # Step 2: Claude from first page
        if not title:
            first_page = extract_first_page_text(path)
            if first_page:
                title = claude_extract_title(rid, fname, first_page)
                method = "claude"
                time.sleep(0.3)  # gentle rate limiting
            else:
                title = old_title
                method = "unchanged"

        results[rid] = title
        changed = "  ✓" if title != old_title else "  ="
        print(f"{changed} [{rid}] ({method}) {title!r}")

    if args.dry_run:
        print(f"\nDry run complete — {len(results)} titles extracted, {skipped} skipped.")
        return

    # Patch umap.json
    updated = 0
    for doc in docs:
        rid = doc["resource_id"]
        if rid in results and results[rid] != doc.get("title"):
            doc["title"] = results[rid]
            updated += 1

    if isinstance(umap, dict):
        umap["docs"] = docs
        out = umap
    else:
        out = docs

    with open(UMAP_PATH, "w") as f:
        json.dump(out, f, indent=2, ensure_ascii=False)

    print(f"\nDone — {updated} titles updated, {skipped} skipped. umap.json written.")
    print("Next: git add app/public/umap.json && git commit && vercel --prod")


if __name__ == "__main__":
    main()
